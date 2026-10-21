"""Local research workbook export. Marketplace text is never an Excel formula."""
import base64
import io
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def safe(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = ''.join(c for c in value if c in '\n\r\t' or ord(c) >= 32)
        if value.lstrip().startswith(('=', '+', '-', '@')):
            value = "'" + value
        return value[:32767]
    return value


def export_workbook(reports, output):
    """Combine expanded reports, single-product comparisons and collection jobs."""
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {}
    headers = {
        '조사 요약': ['자료', '내용'],
        '전체 상품 후보': ['상품군', '플랫폼', '상품 ID', '검색 방식', '상품명', '최저 표시가', '최고 표시가', '통화', '최소 주문량', '판매자', '배송비', '수집 시각', '상품 링크', '검색 출처', '이미지 링크', '검색 이미지 SHA256', '검증 상태', '근거 파일'],
        '가격 비교': ['상품군', '국내 상품', '해외 상품', '국내 표시가', '국내 수량 가정', '국내 개당 원', '해외 개당 원', '차액 원', '국내÷해외 배수', '국내 대비 감소율', '해외 MOQ', '국내 배송비', '매칭 근거', '조건·미확인', '국내 출처', '해외 출처'],
        '사진 대조': ['상품군', '국내 사진', '해외 사진', '관찰 근거', '한계'],
        '검색 실행 기록': ['자료', '플랫폼', '상품군', '방식', '상태', '결과 수', '오류', '근거 파일', 'SHA256'],
        '환율·비용 조건': ['항목', '값', '출처·조건'],
    }
    for name, cols in headers.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        ws.freeze_panes = 'A2'
        sheets[name] = ws

    def append(name, values):
        sheets[name].append([safe(v) for v in values])

    seen = set()
    def candidate(r, source, category=None):
        site = r.get('site', r.get('platform'))
        pid = r.get('product_id', r.get('offer_id'))
        url = r.get('item_url', r.get('url'))
        key = (site, str(pid or url), category or r.get('category'))
        if key in seen:
            return
        seen.add(key)
        append('전체 상품 후보', [category or r.get('category'), site, pid,
            r.get('discovery_methods', r.get('method', r.get('search_type'))),
            r.get('title'), r.get('price_min', r.get('price')), r.get('price_max'),
            r.get('currency'), r.get('moq_value', r.get('moq')), r.get('seller_name', r.get('supplier')),
            r.get('shipping_fee'), r.get('fetched_at', r.get('observed_at')), url,
            r.get('source_url'), r.get('image_url'), r.get('query_image_sha256'),
            '동일 SKU 미확인', r.get('evidence_files', source)])

    for path in reports:
        path = Path(path)
        data = json.loads(path.read_text())
        append('조사 요약', [str(path), data.get('created_at_utc', data.get('action', '저장 조사'))])
        for k, v in data.get('stats', {}).items():
            append('조사 요약', [k, v])
        for limitation in data.get('limitations', []):
            append('환율·비용 조건', ['미확인·한계', limitation, str(path)])
        for r in data.get('candidates', []):
            candidate(r, str(path))
        for j in data.get('jobs', []):
            rows = j.get('products', [])
            for r in rows:
                candidate(r, str(path), j.get('category'))
            append('검색 실행 기록', [str(path), j.get('site'), j.get('category'), j.get('search_type'), j.get('state'), len(rows), j.get('error'), j.get('output_file'), j.get('response_sha256')])
        for r in data.get('runs', []):
            append('검색 실행 기록', [str(path), r.get('site'), r.get('category'), r.get('method'), r.get('quality', r.get('state')), r.get('count'), r.get('error'), r.get('file'), r.get('sha256')])
        for c in data.get('comparisons', []):
            d, o = c['domestic'], c['overseas']
            q = c.get('domestic_unit_count_assumed', 1)
            ws = sheets['가격 비교']
            n = ws.max_row + 1
            append('가격 비교', [c.get('category'), d.get('title'), o.get('title'), d.get('price_min', d.get('price')), q,
                c['domestic_unit_price_krw'], c['overseas_price_krw'], None, None, None,
                o.get('moq_value', o.get('moq')), d.get('shipping_fee'), c.get('match_evidence'),
                c.get('limitations'), d.get('source_url', d.get('url')), o.get('item_url', o.get('url'))])
            # Only our own fixed expressions become formulas; scraped strings stay text.
            ws.cell(n, 8, f'=F{n}-G{n}')
            ws.cell(n, 9, f'=IFERROR(F{n}/G{n},"")')
            ws.cell(n, 10, f'=IFERROR(1-G{n}/F{n},"")').number_format = '0.0%'
            ws.cell(n, 9).number_format = '0.00'
            ps = sheets['사진 대조']
            rowno = ps.max_row + 1
            append('사진 대조', [c.get('category'), None, None, c.get('match_evidence'), c.get('limitations')])
            ps.row_dimensions[rowno].height = 140
            for col, r in [('B', d), ('C', o)]:
                encoded = r.get('inline_image')
                if encoded and encoded.startswith('data:image/'):
                    img = Image(io.BytesIO(base64.b64decode(encoded.split(',', 1)[1])))
                    img.width = img.height = 180
                    ps.add_image(img, f'{col}{rowno}')
        if 'overseas_1688' in data:
            o, d = data['overseas_1688'], data['domestic']
            append('가격 비교', ['창문 필름 (기존 조사)', d['title'], o['selected_option'], d['price'], 1,
                d['price'], o['normalized_price_krw'], data['calculation']['difference_krw'],
                data['calculation']['domestic_price_multiple'], data['calculation']['overseas_cheaper_percent']/100,
                o['moq_m'], d['shipping_fee'], '50cm×2m 환산; 국내 50X200을 cm로 해석',
                '동일 품질·국내 옵션 미확인; 배송/세금 제외', d['url'], o['url']])
            sheets['가격 비교'].cell(sheets['가격 비교'].max_row, 10).number_format = '0.0%'
            append('환율·비용 조건', ['CNY/KRW', data['fx']['cny_krw'], data['fx']['source']])
            for tier in data['alibaba']['tiers_krw']:
                append('환율·비용 조건', ['Alibaba 필름 수량별 가격', tier, '롤 길이 미확인: 면적 비교 불가 · '+data['alibaba']['url']])
    append('조사 요약', ['가격 해석', '후보 노출가와 명시된 수량 가정의 비교. 동일 상품·실결제 원가·수익률 검증 아님.'])
    for ws in wb:
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24445C')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.data_type != 'f' and isinstance(cell.value, str) and cell.value.startswith('https://'):
                    cell.hyperlink = cell.value
                    cell.style = 'Hyperlink'
        for i, cells in enumerate(ws.columns, 1):
            width = max((len(str(c.value or '')) for c in list(cells)[:80]), default=12)
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(15, width * 0.8))
    sheets['사진 대조'].column_dimensions['B'].width = 28
    sheets['사진 대조'].column_dimensions['C'].width = 28
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_name(output.stem + '.tmp.xlsx')
    wb.save(temp)
    temp.replace(output)
    check = load_workbook(output)
    assert list(check.sheetnames) == list(headers)
    check.close()
    return {'output': str(output), 'candidates': len(seen), 'comparisons': sheets['가격 비교'].max_row - 1}
