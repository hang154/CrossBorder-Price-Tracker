"""Verify a terminal paired audit and export its HTML, Excel and monthly joins."""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from app.coupang_purchase_audit import (
    merge_isolated_audit, merge_monthly_opportunities, save_dashboard,
)
from app.sourcing_pipeline import atomic_json
from scripts.coupang_verify_paired_archives import verify


def finalize(run_dir, source, audit_name='virtual-audit'):
    run_dir, source = Path(run_dir), Path(source)
    audit = run_dir / audit_name
    manifest = json.loads((audit / 'manifest.json').read_text())
    if not manifest.get('finished_at_kst'):
        raise ValueError('Audit is still running; no report was rewritten')
    if not manifest.get('profile_removed') or not manifest.get('virtual_display_removed'):
        raise ValueError('Private browser/display cleanup has not been verified')
    queue = json.loads((run_dir / 'queue.json').read_text())
    merge_isolated_audit(run_dir, queue['products'], queue['as_of'], audit_name)
    integrity = verify(audit, run_dir)
    if not integrity['passed']:
        raise ValueError('Archive verification failed: ' + str(integrity['errors']))
    # Validate the source join before replacing the existing reports.
    source_rows = json.loads(source.read_text())['items']
    queue_ids = {str(p['product_id']) for p in queue['products']}
    if any(str(r['domestic']['product_id']) not in queue_ids for r in source_rows):
        raise ValueError('Comparison source contains a product outside this audit')
    save_dashboard(run_dir, queue['products'], queue['as_of'])
    merge_monthly_opportunities(source, run_dir, queue['products'])
    wb = load_workbook(run_dir / 'research.xlsx')

    def sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(["'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v for v in row])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    refresh_file = run_dir / 'ranking-refresh.json'
    if refresh_file.exists():
        refresh = json.loads(refresh_file.read_text())
        rows = []
        for p in refresh['products']:
            o = p.get('observation') or {}
            rows.append([p['product_id'], p.get('item_id'), p.get('vendor_item_id'),
                         p['status'], o.get('title'), o.get('price'), o.get('rating'),
                         o.get('review_count'), o.get('rocket_badge_present'),
                         o.get('delivery_text'), o.get('source_url'), o.get('url')])
        sheet('경로 재확인 상품정보', ['상품 ID', '옵션 ID', '판매 옵션 ID', '상태', '상품명',
              '가격', '별점', '누적 리뷰', '로켓 표시', '배송 문구', '수집 페이지', '원본 상품 URL'], rows)
    sheet('가상 브라우저 시도', ['상품 ID', '상품', '시도 수', '상세 상태', '월 리뷰 상태',
          '최근 1개월 리뷰', '월 판매량 ×15', '마지막 오류'], [
        [p['product_id'], p.get('title'), p.get('isolated_browser_attempts'),
         p.get('detail_status'), p.get('review_status'),
         p.get('monthly_metrics', {}).get('review_count'),
         p.get('estimated_monthly_sales'), p.get('isolated_browser_error')]
        for p in queue['products']])
    wb.save(run_dir / 'research.xlsx')
    atomic_json(run_dir / 'archive-integrity.json', integrity)
    status_file = run_dir / 'route-status.json'
    status = json.loads(status_file.read_text()) if status_file.exists() else {}
    status.update(observed_at_kst=manifest['finished_at_kst'], virtual_browser={
        k: v for k, v in manifest.items() if k not in ('products', 'batches')})
    status['virtual_browser']['manifest'] = audit_name + '/manifest.json'
    status['verdict'] = 'Per-product status is authoritative; attempted does not mean collected.'
    atomic_json(status_file, status)
    complete_reviews = sum(p['review_status'] == 'complete' for p in queue['products'])
    complete_details = sum(p['detail_status'] == 'complete' for p in queue['products'])
    summary = dict(products=len(queue['products']), comparisons=len(source_rows),
                   live_review_collections_complete=complete_reviews,
                   live_detail_archives_complete=complete_details,
                   monthly_sales_nonnull=sum(p.get('estimated_monthly_sales') is not None for p in queue['products']),
                   collection_is_complete=complete_reviews == complete_details == len(queue['products']),
                   archive_integrity=integrity, artifact_sha256={})
    for name in ['index.html', 'research.xlsx', 'queue.json', 'opportunities-with-monthly.json',
                 'route-status.json', 'archive-integrity.json', audit_name + '/manifest.json',
                 audit_name + '/source-queue.json']:
        summary['artifact_sha256'][name] = hashlib.sha256((run_dir / name).read_bytes()).hexdigest()
    atomic_json(run_dir / 'report-verification.json', summary)
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run-dir', type=Path, required=True)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--audit-name', choices=['isolated-audit', 'virtual-audit'], default='virtual-audit')
    args = parser.parse_args()
    result = finalize(args.run_dir, args.source, args.audit_name)
    print(json.dumps({k: v for k, v in result.items() if k != 'artifact_sha256'}, ensure_ascii=False))
    return 0  # Export success is distinct from collection_is_complete.


if __name__ == '__main__':
    raise SystemExit(main())
