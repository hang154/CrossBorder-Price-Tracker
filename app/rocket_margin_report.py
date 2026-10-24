"""Recalculate stored sourcing pairs; no model calls, live crawling or purchases."""
import argparse
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
import hashlib
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from app.rocket_margin import calculate, minimum_price, serializable
from app.bulk_import import prepare_bulk

LABELS=dict(overseas_goods='1688 상품대금',china_inland='중국 내륙배송',sourcing_agent='사입 대행',international_freight='국제운송',insurance='운송보험',customs_duty='관세 가정',other_import_taxes='기타 수입세',import_vat='수입 VAT',clearance='통관 대행',inspection='인증·검수 배분',packaging='포장재',label='라벨·바코드',bundling='세트 구성',center_inbound='센터 입고운송',commission='판매수수료',fulfillment_io='입출고',customer_delivery='고객 배송',storage='보관',return_collection='반품 회수',return_restock='반품 재입고',returned_order_outbound='반품 주문 출고·배송 손실',return_damage_reserve='반품 가치손실 충당',unsold_inventory_reserve='미판매 재고손실 충당',withdrawal='예상 반출',disposal='예상 폐기',ads='광고',fixed_overhead='고정비 배분')
NOTICE=('1688 배대지 합배송 배치 총액을 배분한다. 아래 기존 민감도 총액은 비교용 가정이며 공개 해운요금이나 실제 견적이 아니다. 사업자사입 검증이 완료되기 전 사입추천·확정마진으로 사용하지 않는다. 각 비교 조합은 독립된 가상 배치다. 실제 WING 정산·규격·견적 미확인. 사이즈는 가정이며 공식 요금은 시작 금액이다. 일반과세자의 과세상품 가정으로 VAT 전 현금과 공제 후 손익을 구분한다. 공제0%는 일반과세자의 증빙 미확보 민감도이며 간이과세 계산이 아니다. 원가는 대표 옵션 가격으로 구성·소재·면적이 다른 조합은 재매칭해야 한다. 손익은 반품·재고 손실 충당과 고정비를 반영한 추정치이며 소득세·법인세 전이다. 판매자 예측이 없으면 월 이익·재고 소진·회수일은 비워 둔다. 카탈로그 리뷰×15를 판매자 예측으로 쓰지 않는다.')
SOURCES=[
 {'url':'https://marketplace.coupang.com/rocket-growth','use':'공개 크기별 시작 요금. WING 확정 요금 아님.'},
 {'url':'https://sell.coupang.com/ko-kr/fee-page','use':'판매수수료·물류비 VAT 별도 및 최종 소비자가 기준.'},
 {'url':'https://github.com/gksrudehd123-maker/OMS/blob/main/src/lib/helpers/rg-margin-calc.ts','use':'정산/VAT 분리 검토. 코드 복제 없음. ceil_10은 선택 가능한 참고 반올림이며 공식 정산 규칙 미확인.'},
 {'url':'https://github.com/tramper2/CupangCost/blob/main/src/utils/calculator.ts','use':'BEP·민감도 출력 검토. 자체 요율표를 공식 요율로 채택하지 않음.'},
 {'url':'https://github.com/oiomaker/cowork-plugins-archive/tree/main/moai-commerce/skills','use':'마진·광고 예산·본전 ROAS 판단 개념 검토. 스킬 설치나 광고 집행 없음.'},
 {'url':'https://www.nts.go.kr/nts/cm/cntnts/cntntsView.do?cntntsId=7695&mi=2272','use':'일반과세 매출 VAT − 적격 매입 VAT.'},
 {'url':'https://www.customs.go.kr/kcs/ad/tax/BuyTaxCalculation.do','use':'수입 VAT 과표 관계만 참고. 판매용 수입에 개인직구 면세를 적용하지 않음.'}]

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def safe(v):return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v

def build(source,config,out):
    source,config,out=Path(source),Path(config),Path(out)
    original=json.loads(source.read_text());cfg=json.loads(config.read_text());items=original['items']
    if len({i['pair_id'] for i in items})!=len(items):raise ValueError('Duplicate pair')
    if set(cfg.get('pair_overrides',{}))-{i['pair_id'] for i in items}:raise ValueError('Unknown override pair ID')
    records=[];products=[];audits=[]
    for i in items:
        d,o=i['domestic'],i['overseas'];m=i.get('monthly_review_metrics',{})
        products.append(dict(pair_id=i['pair_id'],title=d['title'],overseas_title=o['title'],group=d['group_name'],listed_price=d['price'],foreign_price=o['price_min'],quantity_verified=i['quantity_verified'],quantity_ratio_assumed=i['quantity_ratio_assumed'],legacy_eligible=i['eligible'],legacy_base_profit=i['scenarios']['base']['estimated_profit'],domestic_url=d['url'],overseas_url=o['item_url'],catalog_monthly_sales=m.get('estimated_monthly_sales'),monthly_status=m.get('status','unavailable'),comparison_issue=i['review']['reason_ko']))
        plan=cfg.get('bulk_plan')
        audit=None
        if plan is not None:
            plan={**plan,**cfg.get('bulk_plan_overrides',{}).get(i['pair_id'],{})}
            audit=prepare_bulk(i,{**cfg['defaults'],**cfg['scenarios'].get('base',{})},plan)
            actual=None
            if audit['ready']:
                actual=serializable(calculate(i,audit['assumptions'],audit['logistics']))
            audits.append(dict(pair_id=i['pair_id'],title=d['title'],moq=o.get('moq_value'),ready=audit['ready'],issues=audit['issues'],metrics=serializable(audit['metrics']),freight=serializable(audit['freight']),actual=actual))
            products[-1]['bulk_ready']=audit['ready']
            products[-1]['bulk_issues']=audit['issues']
        for scenario,s in cfg['scenarios'].items():
            a={**cfg['defaults'],**s,**cfg.get('pair_overrides',{}).get(i['pair_id'],{})}
            for size,logistics in cfg['size_starting_fees_ex_vat'].items():
                calc=calculate(i,a,logistics);nc=calculate(i,a,logistics,credit_ratio=0)
                compact=serializable({k:v for k,v in calc.items() if k!='ledger'})
                ledger=[[l['name'],*serializable([l[k] for k in ['net','vat','deductible_vat','cash']])] for l in calc['ledger']]
                records.append(dict(pair_id=i['pair_id'],scenario=scenario,size=size,
                    minimum_final_price=minimum_price(i,a,logistics),
                    target_final_price=minimum_price(i,a,logistics,target_margin=a['target_margin']),
                    minimum_final_price_no_credit=minimum_price(i,a,logistics,credit_ratio=0),
                    target_final_price_no_credit=minimum_price(i,a,logistics,credit_ratio=0,target_margin=a['target_margin']),
                    profit_no_input_credit=serializable(nc['operating_profit_after_vat']),
                    break_even_roas_no_credit=serializable(nc['break_even_roas']),
                    capital_recovery_days_no_credit=serializable(nc['capital_recovery_days_including_settlement']),
                    **compact,ledger=ledger))
    data=dict(bulk_audits=audits,notice=NOTICE,config=cfg,products=products,records=records,cost_labels=LABELS,input_source_sha256=sha(source),config_sha256=sha(config),built_at_kst=datetime.now(ZoneInfo('Asia/Seoul')).isoformat(),engine_sha256=sha(Path(__file__).with_name('rocket_margin.py')),sources=SOURCES)
    out.mkdir(parents=True,exist_ok=True)
    (out/'report.json').write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')))
    (out/'assumptions.json').write_text(json.dumps(cfg,ensure_ascii=False,indent=2))
    wb=Workbook();ws=wb.active;ws.title='전체 마진 민감도'
    headers=['조합ID','국내 상품명','비용 가정','크기 가정','최종 소비자가','해외 표시가','기존 기준 공헌이익','VAT 전 현금기여(충당 포함)','배분 VAT 납부액','공제·고정비 후 이익','소비자가 대비 마진','광고 전 이익','본전 ROAS','최저 최종판매가','목표마진 최종판매가','공제0% 이익','초기 재고현금','재고소진일','원금회수일(정산지연 제외)','원금회수일(정산지연 포함)','판매자 월이익','카탈로그 월추정(판매자예측 아님)','월 집계상태','구성 수량근거','국내 URL','해외 URL']
    ws.append(headers);by={p['pair_id']:p for p in products}
    keys=['cash_contribution_before_vat','vat_due_allocated','operating_profit_after_vat','margin_on_consumer_price','profit_before_ads_after_vat','break_even_roas','minimum_final_price','target_final_price','profit_no_input_credit','initial_inventory_cash','inventory_days','capital_recovery_days_excluding_settlement','capital_recovery_days_including_settlement','monthly_operating_profit']
    for r in records:
        p=by[r['pair_id']];values=[r['pair_id'],p['title'],r['scenario'],r['size'],r['final_consumer_price'],p['foreign_price'],p['legacy_base_profit'],*[r[k] for k in keys],p['catalog_monthly_sales'],p['monthly_status'],p['quantity_verified'],p['domestic_url'],p['overseas_url']]
        assert len(values)==len(headers)
        ws.append([safe(v) for v in values])
    detail=wb.create_sheet('기준 중형 원가내역');detail.append(['조합ID','상품','비용 항목','공급가/충당','VAT','공제 가정액','현금/충당액'])
    for r in records:
        if r['scenario']=='base' and r['size']=='중형':
            for name,net,vat,credit,cash in r['ledger']:detail.append([safe(r['pair_id']),safe(by[r['pair_id']]['title']),LABELS[name],net,vat,credit,cash])
    allocation=wb.create_sheet('배대지 합배송 배분')
    allocation.append(['조합ID','상품','비용 가정','배분 기준','견적 상태','배치 실물 수량','해당 SKU 실물 수량','해당 SKU 판매단위 수량','비용 항목','배치 총액','해당 SKU 배분액','판매단위당 배분액'])
    for r in records:
        sh=r.get('shipment_allocation')
        if sh and r['size']==next(iter(cfg['size_starting_fees_ex_vat'])):
            for cost,total in sh['shipment_totals'].items():
                allocation.append([safe(r['pair_id']),safe(by[r['pair_id']]['title']),r['scenario'],sh['basis'],sh['quote_status'],sh['shipment_physical_units'],sh['sku_physical_units'],sh['sku_selling_units'],LABELS.get(cost,cost),total,sh['sku_allocated_totals'][cost],sh['per_sale_costs'][cost]])
    audit_sheet=wb.create_sheet('사업자사입 검증')
    audit_sheet.append(['조합ID','상품','표시 MOQ','실제 원가 계산 준비','누락·불일치','실제 VAT 후 이익'])
    for audit in audits:
        audit_sheet.append([safe(audit['pair_id']),safe(audit['title']),audit['moq'],audit['ready'],safe('; '.join(audit['issues'])),audit['actual']['operating_profit_after_vat'] if audit['actual'] else None])
    tariff_sheet=wb.create_sheet('해운 노선 요금')
    tariff_sheet.append(['노선','요금 및 경계'])
    for route,tariff in cfg.get('shipping_tariffs',{}).items():tariff_sheet.append([route,safe(json.dumps(tariff,ensure_ascii=False))])
    for name,obj in [('가정과 한계',{'notice':NOTICE,**cfg}),('근거',data['sources'])]:
        sheet=wb.create_sheet(name);sheet.append(['항목','내용'])
        for k,v in (obj.items() if isinstance(obj,dict) else enumerate(obj)):sheet.append([str(k),safe(json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v)])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='243746')
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(55,max(16,len(str(col[0].value))+3))
    wb.save(out/'로켓그로스_마진분석.xlsx')
    template=Path(__file__).with_name('rocket_margin.html').read_text()
    (out/'index.html').write_text(template.replace('__REPORT_DATA__',json.dumps(data,ensure_ascii=False,separators=(',',':')).replace('<','\\u003c')))
    baseline=[r for r in records if r['scenario']=='base' and r['size']=='중형'];eligible={p['pair_id'] for p in products if p['legacy_eligible']}
    summary=dict(bulk_ready=sum(x['ready'] for x in audits),bulk_pending=sum(not x['ready'] for x in audits),pairs=len(items),scenario_size_combinations=len(records),baseline='base costs / assumed medium package / conditional full VAT credit',baseline_positive=sum(r['operating_profit_after_vat']>0 for r in baseline),baseline_positive_legacy_candidates=sum(r['operating_profit_after_vat']>0 and r['pair_id'] in eligible for r in baseline),seller_forecast_supplied=any(r['monthly_operating_profit'] is not None for r in records),ledger_reconciliations=len(records)*2,artifacts={f:sha(out/f) for f in ['report.json','index.html','로켓그로스_마진분석.xlsx','assumptions.json']})
    (out/'verification.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2));return summary

def main():
    p=argparse.ArgumentParser();p.add_argument('--source',required=True,type=Path);p.add_argument('--config',required=True,type=Path);p.add_argument('--out',required=True,type=Path);args=p.parse_args()
    print(json.dumps(build(args.source,args.config,args.out),ensure_ascii=False,indent=2))

if __name__=='__main__':main()
