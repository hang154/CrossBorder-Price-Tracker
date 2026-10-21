from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.template_excel_agent import write_workbook, slots, TemplateAgent


def template(tmp_path):
    p=tmp_path/'original.xlsx'; w=Workbook();s=w.active
    s.merge_cells('B2:D2');s['B2']='기본 정보';s['B4']=1;s['C4']='기존 상품';s['M4']='=J4*0.89'
    for r in range(5,8):
        s.cell(r,2,r-3);s.cell(r,8,f'=G{r}*15');s.cell(r,21,f'=T{r}*300')
    w.save(p);return p


def inputs():
    rows=[dict(keyword='호안석 팔찌',season='연중(가설)',domestic_id='coupang/1',foreign_id='1688/2',wholesale_id='',comparison_note='동일 상품 미확정')]
    sources={'coupang/1':dict(key='coupang/1',url='https://www.coupang.com/vp/products/1',title='호안석 팔찌',price=20000,review_count=9999),
             '1688/2':dict(key='1688/2',url='https://detail.1688.com/offer/2.html',title='호안석 팔찌',currency='KRW',price=2500)}
    return rows,sources


def export(tmp_path, mutate=None):
    src=template(tmp_path);original=src.read_bytes();rows,sources=inputs()
    if mutate:mutate(rows,sources)
    target=tmp_path/'filled.xlsx'
    managed=write_workbook(src,target,rows,sources,dict(fee_rate=.11,logistics_krw=3000,krw_per_cny=300),{})
    assert src.read_bytes()==original
    return src,load_workbook(target),managed


def test_preserves_template_and_existing_cells(tmp_path):
    src,w,managed=export(tmp_path);s=w.active
    assert managed==[5];assert s['C4'].value=='기존 상품';assert s['M4'].value=='=J4*0.89'
    assert str(s.merged_cells)==str(load_workbook(src).active.merged_cells)
    assert s['H6'].value=='=G6*15'
    assert w.sheetnames[1:]==['조사근거','에이전트설정','실행검증']


def test_krw_not_fabricated_cny_and_no_cumulative_sales(tmp_path):
    _,w,_=export(tmp_path);s=w.active
    assert s['T5'].value is None
    assert s['U5'].value.startswith("='조사근거'!B")
    assert s['H5'].value is None
    assert 'COUNT(R5,W5)=0' in s['N5'].value
    assert s['O5'].value is None and s['P5'].value is None


def test_complete_dated_reviews_can_populate_zero(tmp_path):
    def add(rows,sources):
        sources['coupang/1']['review_audit']={'windows':{'1':dict(start_inclusive='2026-09-21',end_inclusive='2026-10-21',is_complete=True,estimated_sales=0)}}
    _,w,_=export(tmp_path,add)
    assert w.active['H5'].value==0
    assert w.active['G5'].value=='2026-09-21 ~ 2026-10-21'


def test_website_formula_injection_escaped(tmp_path):
    _,w,_=export(tmp_path,lambda rows,sources:rows[0].update(keyword='=WEBSERVICE("evil")'))
    assert w.active['C5'].data_type=='s'
    assert w.active['C5'].value.startswith("'=")
    assert w.active['K5'].data_type=='f'


def test_stage_rejects_unobserved_source_and_nonexact_quote():
    a=object.__new__(TemplateAgent);a.count=1;a.state={'rows':[],'sources':{}}
    with pytest.raises(ValueError,match='observed'):
        a.dispatch('stage_row',dict(keyword='a',domestic_id='coupang/unknown',domestic_quote='a',foreign_id='',foreign_quote='',wholesale_id='',wholesale_quote=''))


def test_formula_caches_show_known_calculations(tmp_path):
    export(tmp_path)
    w=load_workbook(tmp_path/'filled.xlsx',data_only=True);s=w.active
    assert s['U5'].value==2500
    assert s['M5'].value==14800
    assert s['V5'].value==12300
    assert s['N5'].value==pytest.approx(.615)
    assert s['X5'].value==pytest.approx(1/.615)
    assert s['I5'].value is None


def test_missing_both_costs_has_no_margin_or_roas(tmp_path):
    export(tmp_path,lambda rows,sources: rows[0].update(foreign_id=''))
    w=load_workbook(tmp_path/'filled.xlsx',data_only=True)
    assert w.active['N5'].value is None
    assert w.active['X5'].value is None
