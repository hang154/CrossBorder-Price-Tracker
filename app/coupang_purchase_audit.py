"""Resume paired one-month review and detail collection for source-pinned products."""
import argparse
from concurrent.futures import ThreadPoolExecutor
import datetime as dt
import hashlib
import html
import io
import json
import time
from pathlib import Path
from urllib.parse import urlsplit
import httpx
from PIL import Image
from openpyxl import Workbook
from app.coupang_reviews import crawl, CollectionCoordinator, Browser, KST
from app.coupang_navigation import home_search_detail
from app.sourcing_pipeline import atomic_json
from scripts.coupang_detail_collect import EXTRACT, SafeHTML, document


def monthly_metrics(report):
 w=report['windows']['1']; c=report.get('coverage',{});observed=c.get('pages',0)>0
 covered=bool(c.get('date_range_complete') or c.get('covered_months',0)>=1) and report.get('missing_dates',0)==0
 count=w['rating_or_review_events'] if observed else None
 complete=covered and c.get('rating_only_coverage')=='count_reconciled'
 return dict(start_inclusive=w['start_inclusive'],end_inclusive=w['end_inclusive'],date_basis=report['date_basis'],review_count=count,
 text_media_reviews=w['review_events'] if observed else None,rating_only_records=w['rating_only_events'] if observed else None,
 estimated_monthly_sales=count*15 if complete and count is not None else None,
 observed_review_sales_estimate=count*15 if count is not None else None,
 formula='distinct dated review_id records in rolling calendar month × 15; not unique reviewers',multiplier=15,
 date_range_covered=covered,all_rating_coverage_reconciled=complete,status='complete' if complete else 'partial' if observed else 'unavailable')


def capture_detail(browser,product,root):
 folder=root/str(product);folder.mkdir(parents=True,exist_ok=True)
 try:
  more=browser.call('eval',"[...document.querySelectorAll('button')].some(e=>e.innerText.trim()==='상품정보 더보기')")
  if more:browser.call('click','--role','button','--name','상품정보 더보기')
  data=browser.call('eval',EXTRACT)
  # React can keep the collapsed DOM briefly after the native click.
  for _ in range(20):
   if data.get('expanded') or data.get('error') or str(data.get('product_id'))!=str(product):break
   time.sleep(.2);data=browser.call('eval',EXTRACT)
  if data.get('error') or str(data.get('product_id'))!=str(product):raise ValueError('Detail product identity/content missing')
  raw=data.pop('detail_html');(folder/'detail-original.html.txt').write_text(raw);(folder/'description.txt').write_text(data.get('detail_text') or '')
  data.update(observed_at_kst=dt.datetime.now(KST).isoformat(),raw_html_sha256=hashlib.sha256(raw.encode()).hexdigest(),session_kind='paired_review_detail_same_visit',limitations=['선택된 옵션 기준이며 모든 옵션을 개별 방문한 자료는 아닙니다.','이미지/설명의 판매자 주장은 검증되지 않았습니다.'])
  # Save observed content before downloading assets, including on a later interruption.
  data['status']='observed_assets_pending';atomic_json(folder/'detail.json',data)
  entries=data.get('detail_images',[])+([data['main_image']] if data.get('main_image') else [])
  entries=list({e['src']:e for e in entries if e.get('src')}.values())
  def asset(pair):
   i,e=pair;original=e['src'];url='https:'+original if original.startswith('//') else original;host=urlsplit(url).hostname or ''
   a=dict(source_url=original)
   if urlsplit(url).scheme not in ('https','http') or not host.endswith('.coupangcdn.com'):return {**a,'error':'unsupported_origin'}
   try:
    r=httpx.get(url,timeout=35,follow_redirects=True);r.raise_for_status();b=r.content
    with Image.open(io.BytesIO(b)) as im:width,height=im.size;fmt=im.format;im.verify()
    ext={'JPEG':'jpg','PNG':'png','WEBP':'webp','GIF':'gif','AVIF':'avif'}.get(fmt,'img');file=f'images/{i:03d}.{ext}';(folder/'images').mkdir(exist_ok=True);(folder/file).write_bytes(b)
    return {**a,'file':file,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'width':width,'height':height}
   except Exception as exc:return {**a,'error':str(exc)}
  with ThreadPoolExecutor(max_workers=4) as pool:assets=list(pool.map(asset,enumerate(entries,1)))
  data['assets']=assets;local={a['source_url']:a['file'] for a in assets if a.get('file')}
  data['status']='complete' if entries and data.get('expanded') and len(local)==len(entries) else 'partial'
  safe=SafeHTML(local);safe.feed(raw);safe.close();atomic_json(folder/'detail.json',data)
  (folder/'index.html').write_text(document(data.get('title',str(product)),'<h1>'+html.escape(data.get('title') or str(product))+'</h1><p>'+data['status']+'</p>'+''.join(safe.parts)))
  return data
 except Exception as exc:
  data=dict(product_id=str(product),status='failed',error=str(exc));atomic_json(folder/'detail.json',data);return data


def queue_from_results(path):
 src=json.loads(path.read_text());by={}
 for r in src['items']:
  d=dict(r['domestic']);pid=str(d['product_id'])
  value=dict(d,opportunity_score=r['opportunity']['overall_opportunity_score'],monthly_review_count=None,estimated_monthly_sales=None,review_status='pending',detail_status='pending')
  if pid not in by or value['opportunity_score']>by[pid]['opportunity_score']:by[pid]=value
 audit=path.parent/'shortlist-evidence-audit.json'
 priority=[]
 if audit.exists():priority=[x.split(':')[1] for x in json.loads(audit.read_text())['retained_for_quote_validation_only']]
 return sorted(by.values(),key=lambda r:(str(r['product_id']) not in priority,-r['opportunity_score']))


def merge_isolated_audit(out,queue,as_of,audit_name='isolated-audit'):
 """Merge only terminal results for the exact frozen queue and date."""
 if audit_name not in ('isolated-audit','virtual-audit'):raise ValueError('Unsupported audit directory')
 path=out/audit_name/'manifest.json'
 state=json.loads(path.read_text())
 if not state.get('finished_at_kst'):raise ValueError('Isolated audit has not finished')
 if state.get('as_of')!=as_of:raise ValueError('Isolated audit date mismatch')
 raw=(out/'queue.json').read_bytes();snapshot=out/audit_name/'source-queue.json'
 if state.get('source_sha256')==hashlib.sha256(raw).hexdigest():
  if not snapshot.exists():snapshot.write_bytes(raw)
 elif snapshot.exists() and hashlib.sha256(snapshot.read_bytes()).hexdigest()==state.get('source_sha256'):
  frozen=json.loads(snapshot.read_text());original={str(p['product_id']):p for p in frozen['products']}
  mutable={'review_status','detail_status','monthly_metrics','monthly_review_count','estimated_monthly_sales','isolated_browser_attempts','isolated_browser_error','review_evidence_root','detail_evidence_root'}
  if set(original)!={str(p['product_id']) for p in queue}:raise ValueError('Isolated audit source queue mismatch')
  for p in queue:
   if any(p.get(k)!=v for k,v in original[str(p['product_id'])].items() if k not in mutable):raise ValueError('Isolated audit source queue mismatch')
 else:raise ValueError('Isolated audit source queue mismatch')
 known={str(p['product_id']) for p in queue}
 if not set(state['products']).issubset(known):raise ValueError('Unknown isolated audit product')
 for p in queue:
  entry=state['products'].get(str(p['product_id']))
  if entry is None:continue
  p['isolated_browser_attempts']=len(entry.get('attempts',[]))
  p['isolated_browser_error']=next((a.get('error') for a in reversed(entry.get('attempts',[])) if a.get('error')),None)
  if p.get('review_status')!='complete' and entry.get('monthly_metrics'):
   metric=entry['monthly_metrics'];p.update(monthly_metrics=metric,review_status=metric['status'],monthly_review_count=metric.get('review_count'),estimated_monthly_sales=metric.get('estimated_monthly_sales'))
   if metric['status'] in ('complete','partial'):p['review_evidence_root']=audit_name
  if p.get('detail_status')!='complete':
   status=entry.get('detail_status','pending');p['detail_status']='not_reached' if status=='pending' else status
   if status in ('complete','partial'):p['detail_evidence_root']=audit_name
 return state


def save_dashboard(out,queue,as_of):
 wb=Workbook();ws=wb.active;ws.title='월 리뷰 판매량·상세 현황'
 headers=['상품ID','상품','상품군','리뷰 상태','최근 1개월 리뷰','×15 추정판매량(기간·범위확인)','부분 관측 ×15','상세 상태','가격','별점','누적 리뷰','로켓 표시','원본 URL']
 ws.append(headers);rows=[]
 for p in queue:
  metric=p.get('monthly_metrics',{});values=[p['product_id'],p.get('title'),p.get('group_name'),p['review_status'],metric.get('review_count'),metric.get('estimated_monthly_sales'),metric.get('observed_review_sales_estimate'),p['detail_status'],p.get('price'),p.get('rating'),p.get('review_count'),p.get('rocket_badge_present'),p.get('url')]
  ws.append(["'"+x if isinstance(x,str) and x.startswith(('=','+','-','@')) else x for x in values])
  pid=str(p['product_id']);links=[]
  for kind,label,key in [('reviews','리뷰 근거','review_evidence_root'),('details','상세 보기','detail_evidence_root')]:
   relative=Path(p.get(key,''))/kind/pid/'index.html'
   if (out/relative).exists():links.append('<a href="'+html.escape(relative.as_posix(),quote=True)+'">'+label+'</a>')
  if p.get('url') and urlsplit(p['url']).scheme in ('http','https'):links.append('<a href="'+html.escape(p['url'],quote=True)+'">원본 상품</a>')
  rows.append('<tr>'+''.join('<td>'+html.escape(str(x) if x is not None else '미확보')+'</td>' for x in values[:-1])+'<td>'+' · '.join(links)+'</td></tr>')
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(out/'research.xlsx');atomic_json(out/'queue.json',{'as_of':as_of,'products':queue})
 body='<h1>최근 1개월 리뷰 ×15 · 상품 상세 전수 조사</h1><p>기준일 '+html.escape(as_of)+' · 리뷰 ID 중복 제거, 한국 시간 기준 1개 역월(양끝 날짜 포함). 참여자 수와 구분합니다.</p><p>미확보는 0이 아닙니다. 부분 관측 추정은 전체 월 판매량으로 사용하지 않습니다. 누적 리뷰를 월 리뷰로 바꾸지 않습니다.</p><p><a href="research.xlsx">엑셀</a> · <a href="route-status.json">실제 경로 관측</a></p><table><thead><tr>'+''.join('<th>'+h+'</th>' for h in headers[:-1]+['근거'])+'</tr></thead><tbody>'+''.join(rows)+'</tbody></table>'
 live=''
 rotation=out/'full-rotation'/'manifest.json'
 if rotation.exists():
  r=json.loads(rotation.read_text())
  live+='<p>HTTP 경로 순회: '+str(r.get('processed_this_round',0))+'/'+str(r.get('products',0))+'개 · 요청 '+str(r.get('requests_this_execution',0))+'회 · 유효 단계 응답 '+str(r.get('observed_stages',0))+'건 · '+('종료' if r.get('finished_at_kst') else '진행 중')+' · <a href="full-rotation/manifest.json">상품별 결과</a></p>'
 refresh=out/'ranking-refresh.json'
 if refresh.exists():
  r=json.loads(refresh.read_text());live+='<p>랭킹 경로 재수집: '+str(r['source_pages_observed'])+'/'+str(r['source_pages_requested'])+'개 분류 페이지 · 상품/옵션 일치 '+str(r['matched_products'])+'개. 가격·별점·누적 리뷰·배송 표시 자료이며, 월 리뷰/완성 상세와 구분합니다. <a href="ranking-refresh.json">수집 원본</a></p>'
 for audit_name,label in [('isolated-audit','이전 Chrome'),('virtual-audit','가상 화면 Chrome')]:
  isolated=out/audit_name/'manifest.json'
  if isolated.exists():
   r=json.loads(isolated.read_text());live+='<p>'+label+' 홈→검색→상품 링크: '+str(r.get('processed',0))+'/'+str(r.get('total',0))+'개 · 상세 완료 '+str(r.get('detail_complete',0))+'개 · 월 리뷰 완료 '+str(r.get('reviews_complete',0))+'개 · 월 리뷰 부분 관측 '+str(r.get('reviews_partial',0))+'개 · '+('종료' if r.get('finished_at_kst') else '진행 중')+' · <a href="'+audit_name+'/manifest.json">상품별 시도·오류</a></p>'
   browser_errors=sum('NoSuchWindowException' in str(a.get('error')) or 'InvalidSessionIdException' in str(a.get('error')) for p in r.get('products',{}).values() for a in p.get('attempts',[]))
   if browser_errors:live+='<p>이 실행에는 브라우저 세션 오류 '+str(browser_errors)+'회가 포함됩니다. 사이트 접속 차단과 구분하며, 모든 상품이 사이트까지 도달했다는 뜻은 아닙니다.</p>'
 body=body.replace('<table>',live+'<table>',1)
 (out/'index.html').write_text(document('월 리뷰·상세 조사',body))


def merge_monthly_opportunities(source,out,queue):
 original=json.loads(source.read_text());metrics={str(p['product_id']):p.get('monthly_metrics') for p in queue}
 result=[]
 for row in original['items']:
  r=dict(row);m=metrics.get(str(r['domestic']['product_id']))
  r['monthly_review_metrics']=m
  r['estimated_monthly_sales']=m.get('estimated_monthly_sales') if m else None
  r['conditional_monthly_contribution']={name:round(v['estimated_profit']*r['estimated_monthly_sales'],2) if r['estimated_monthly_sales'] is not None else None for name,v in r['scenarios'].items()}
  r['monthly_estimate_scope']='Coupang product catalog, potentially multiple sellers/options; not forecast of this seller sales'
  result.append(r)
 result.sort(key=lambda r:(bool(r.get('eligible')),r['estimated_monthly_sales'] is not None,r['conditional_monthly_contribution'].get('base') if r['conditional_monthly_contribution'].get('base') is not None else float('-inf'),r.get('opportunity',{}).get('overall_opportunity_score',0)),reverse=True)
 atomic_json(out/'opportunities-with-monthly.json',{'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'formula':'dated distinct review IDs in one calendar month × 15','items':result})
 from openpyxl import load_workbook
 wb=load_workbook(out/'research.xlsx');ws=wb.create_sheet('조합별 월 추정 수익');ws.append(['비교ID','국내 상품','월 리뷰','추정 월 판매량','기준 시나리오 주문이익','조건부 월 공헌이익','월 리뷰 상태','단위 확정','원본 URL'])
 for r in result:
  m=r['monthly_review_metrics'] or {};ws.append([r['pair_id'],r['domestic']['title'],m.get('review_count'),r['estimated_monthly_sales'],r['scenarios']['base']['estimated_profit'],r['conditional_monthly_contribution']['base'],m.get('status','pending'),r['quantity_verified'],r['domestic']['url']])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(out/'research.xlsx')


def main():
 ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,required=True);ap.add_argument('--out',type=Path,required=True);ap.add_argument('--opencli-main',type=Path,required=True);ap.add_argument('--max',type=int,default=137);ap.add_argument('--max-pages',type=int,default=100);args=ap.parse_args()
 out=args.out;out.mkdir(parents=True,exist_ok=True);as_of=dt.datetime.now(KST).date()
 queue=queue_from_results(args.source)
 if (out/'queue.json').exists():
  cached=json.loads((out/'queue.json').read_text())
  if cached['as_of']!=as_of.isoformat():raise ValueError('Date changed: use a fresh output directory for a fresh rolling month')
  prior={str(r['product_id']):r for r in cached['products']};queue=[prior.get(str(r['product_id']),r) for r in queue]
 save_dashboard(out,queue,as_of.isoformat())
 coordinator=CollectionCoordinator(0,'background');prefix=['node',str(args.opencli_main.resolve())]
 for p in queue[:args.max]:
  if p['review_status']=='complete' and p['detail_status']=='complete':continue
  if coordinator.stopped.is_set():break
  pid=str(p['product_id'])
  if p['review_status']=='complete' and p['detail_status']!='complete':
   b=Browser(prefix,'coupang-detail-resume-'+pid,coordinator);trace=[]
   try:
    home_search_detail(b,pid,p.get('title') or p.get('group_name'),trace)
    p['detail_status']=capture_detail(b,pid,out/'details')['status']
   except Exception as exc:
    p['detail_status']='failed';p['detail_error']=str(exc)
   finally:
    atomic_json(out/'details'/pid/'navigation.json',trace)
    try:b.call('close')
    except Exception:pass
   save_dashboard(out,queue,as_of.isoformat());continue
  p['review_status']='running';p['detail_status']='pending';save_dashboard(out,queue,as_of.isoformat())
  def observer(browser,product):
   data=capture_detail(browser,product,out/'details');p['detail_status']=data['status']
  report=crawl('https://www.coupang.com/vp/products/'+pid,out/'reviews'/pid,prefix,as_of,max_pages=args.max_pages,resume=(out/'reviews'/pid/'state.json').exists(),coordinator=coordinator,navigation_query=p.get('title') or p.get('group_name'),detail_observer=observer,months_to_cover=1)
  metric=monthly_metrics(report);p['monthly_metrics']=metric;p['review_status']=metric['status'];p['monthly_review_count']=metric['review_count'];p['estimated_monthly_sales']=metric['estimated_monthly_sales']
  if p['detail_status']=='pending':p['detail_status']='not_reached'
  p['review_error']=report.get('coverage',{}).get('error');save_dashboard(out,queue,as_of.isoformat())
 atomic_json(out/'summary.json',dict(total=len(queue),reviews_complete=sum(p['review_status']=='complete' for p in queue),reviews_partial=sum(p['review_status']=='partial' for p in queue),reviews_unavailable=sum(p['review_status']=='unavailable' for p in queue),details_complete=sum(p['detail_status']=='complete' for p in queue),pending=sum(p['review_status']=='pending' for p in queue),stopped_on_access_restriction=coordinator.stopped.is_set(),multiplier=15,source_sha256=hashlib.sha256(args.source.read_bytes()).hexdigest(),finished_at_kst=dt.datetime.now(KST).isoformat()))
 merge_monthly_opportunities(args.source,out,queue)
 print(json.dumps(json.loads((out/'summary.json').read_text()),ensure_ascii=False),flush=True)
 return 1 if coordinator.stopped.is_set() else 0
if __name__=='__main__':raise SystemExit(main())
