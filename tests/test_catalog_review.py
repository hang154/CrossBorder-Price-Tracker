import hashlib
import json

from openpyxl import load_workbook

from app.catalog_review import records, report, seller_signals
from app import catalog_review


def fixture(tmp_path):
    source = tmp_path / 'raw.json'
    source.write_text('{}')
    row = {'product_id': '123', 'title': '</script><img onerror=alert(1)>',
           'item_url': 'https://www.alibaba.com/product-detail/Test_123.html',
           'price_min': 2, 'currency': 'USD', 'seller_name': '=HYPERLINK("bad")',
           'evidence_text': '7 years · 100 sold · 4.8/5'}
    jobs = [{'site': 'alibaba', 'topic_id': topic, 'state': 'collected', 'rows': [row],
             'file': str(source), 'sha256': hashlib.sha256(source.read_bytes()).hexdigest()}
            for topic in ['a', 'b']]
    topics = [{'id': topic, 'material': 'stone', 'label_ko': topic, 'form': 'bracelet',
               'quality_target': 'unverified', 'query_zh': '石', 'query_en': 'stone'} for topic in ['a', 'b']]
    (tmp_path / 'catalog.json').write_text(json.dumps({'topics': topics, 'jobs': jobs}))
    (tmp_path / 'annotations.json').write_text('{"items":[]}')


def test_dedup_retains_both_source_observations(tmp_path):
    fixture(tmp_path)
    rows, _, _ = records(tmp_path)
    assert len(rows) == 1
    assert rows[0]['topics'] == ['a', 'b']
    assert len(rows[0]['observations']) == 2


def test_metrics_are_observed_quotes_not_trust_scores():
    signals = seller_signals('7 years · 100+ 구매 · 재구매율 23% · 4.8/5')
    assert signals['repurchase_rate'] == ['재구매율 23%']
    assert signals['rating'] == ['4.8/5']
    assert 'trust_score' not in signals
    assert seller_signals('natural AAA stone') == {}


def test_report_escapes_script_and_workbook_formula(tmp_path):
    fixture(tmp_path)
    report(tmp_path)
    page = (tmp_path / 'index.html').read_text()
    data = page.split('<script id="data" type="application/json">')[1].split('</script>')[0]
    assert json.loads(data)[0]['title'].startswith('</script>')
    assert '</script>' not in data
    workbook = load_workbook(tmp_path / 'catalog.xlsx')
    cell = workbook['상품 후보'].cell(2, 14)
    assert cell.data_type == 's'
    assert cell.value.startswith("'=")


def test_resume_requests_only_missing_candidates_and_rejects_invented_quotes(tmp_path, monkeypatch):
    fixture(tmp_path)
    path = tmp_path/'catalog.json'
    source = json.loads(path.read_text())
    source['jobs'][0]['rows'].append({**source['jobs'][0]['rows'][0], 'product_id': '456', 'title': 'stone bracelet'})
    path.write_text(json.dumps(source))
    previous = {'key': 'alibaba/123', 'reason_ko': 'preserved prior review'}
    (tmp_path/'annotations.json').write_text(json.dumps({'items': [previous]}))
    calls = []
    def fake_call(out, name, prompt, schema):
        calls.append(prompt)
        return {'items': [{'key':'alibaba/456','relevance':'relevant','form':'bracelet',
            'material_quote':'stone','quality_quote':'certified AAA',
            'quality_claim_tier':'standard_or_unspecified','reason_ko':'unverified grade'}]}
    monkeypatch.setattr(catalog_review, 'qwen_call', fake_call)
    catalog_review.classify(tmp_path)
    assert len(calls) == 1
    assert 'alibaba/123' not in calls[0]
    items = json.loads((tmp_path/'annotations.json').read_text())['items']
    assert items[0] == previous
    assert items[1]['quality_quote'] == ''
    assert items[1]['material_quote'] == 'stone'
