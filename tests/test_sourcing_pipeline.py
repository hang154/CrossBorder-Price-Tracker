import json
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook
from app.sourcing_pipeline import compare_pair, run_pipeline
from app.sourcing_excel import export_workbook


def rows():
    return [dict(platform='naver', product_id='1', price=11990, currency='KRW', title='Domestic'),
            dict(platform='alibaba', product_id='2', price=975, currency='KRW', title='Overseas')]


def pair():
    return dict(category='oil', domestic=dict(platform='naver', product_id='1', units=2),
                overseas=dict(platform='alibaba', product_id='2', units=1),
                match_evidence='Image shows 1+1', limitations='Bundle not verified')


def test_bundle_normalization_and_denominators():
    c = compare_pair(pair(), rows(), {})
    assert c['domestic_unit_price_krw'] == 5995
    assert c['difference_krw'] == 5020
    assert c['domestic_multiple'] == pytest.approx(5995/975)


@pytest.mark.parametrize('units', [0, -1, float('nan')])
def test_invalid_units_cannot_create_price_advantage(units):
    p=pair();p['overseas']['units']=units
    with pytest.raises(ValueError): compare_pair(p, rows(), {})


def test_unknown_fx_and_missing_matching_evidence_rejected():
    r=rows();r[1]['currency']='CNY'
    with pytest.raises(ValueError, match='exchange rate'):compare_pair(pair(), r, {})
    p=pair();p['match_evidence']=''
    with pytest.raises(ValueError, match='evidence'):compare_pair(p, rows(), {})


def test_workbook_keeps_marketplace_formulas_as_text_and_own_formulas_live(tmp_path):
    c=compare_pair(pair(), rows(), {})
    data={'candidates':[dict(site='naver',product_id='1',title='=HYPERLINK("evil")',price_min=11990)],'comparisons':[c]}
    p=tmp_path/'data.json';p.write_text(json.dumps(data))
    f=tmp_path/'research.xlsx';export_workbook([p],f)
    w=load_workbook(f)
    assert w['전체 상품 후보']['E2'].data_type=='s'
    assert w['전체 상품 후보']['E2'].value.startswith("'=")
    assert w['가격 비교']['H2'].value=='=F2-G2'
    assert w['가격 비교']['E2'].value==2
    assert len(w.sheetnames)==6


def test_resume_preserves_success_and_retries_only_failed_jobs(tmp_path, monkeypatch):
    config=tmp_path/'config.json';config.write_text(json.dumps({'products':[dict(id='oil',query_ko='oil',query_en='oil')]}))
    monkeypatch.setattr('app.sourcing_pipeline.browser_ready',lambda _:dict(ready=True))
    calls=[]
    def run(argv, **kwargs):
        calls.append(argv)
        if 'alibaba' in argv and len([c for c in calls if 'alibaba' in c])==1:
            return SimpleNamespace(returncode=1,stdout=json.dumps({'ok':False,'error':{'message':'temporary'}}))
        n='naver-shopping' in argv
        row=dict(product_id='123',item_url='https://cr3.shopping.naver.com/v2/bridge/searchGate?nv_mid=123' if n else 'https://www.alibaba.com/product-detail/Test_123.html',source_url='https://search.naver.com/search.naver?query=oil',price_min=100,currency='KRW')
        return SimpleNamespace(returncode=0,stdout=json.dumps({'ok':True,'data':[row]}))
    monkeypatch.setattr('app.sourcing_pipeline.subprocess.run',run)
    out=tmp_path/'out'
    assert run_pipeline(config,out,['opencli'])['complete'] is False
    assert len(calls)==2
    assert run_pipeline(config,out,['opencli'],resume=True,retry_failed=True)['complete'] is True
    assert len(calls)==3 and 'alibaba' in calls[-1]
    state=json.loads((out/'report.json').read_text())
    assert state['jobs'][1]['previous_attempts'][0]['state']=='failed'
    next(out.glob('*.json')).touch()
    (out/next(j['key']+'.json' for j in state['jobs'] if j['state']=='collected')).write_text('changed')
    with pytest.raises(ValueError,match='response missing/changed'):run_pipeline(config,out,['opencli'],resume=True)


def test_changed_configuration_rejected_on_resume(tmp_path,monkeypatch):
    monkeypatch.setattr('app.sourcing_pipeline.browser_ready',lambda _:dict(ready=True))
    p=tmp_path/'c.json';p.write_text('{"products":[]}')
    run_pipeline(p,tmp_path/'out',['opencli'])
    p.write_text('{"products":[],"limit":7}')
    with pytest.raises(ValueError,match='changed'):run_pipeline(p,tmp_path/'out',['opencli'],resume=True)


def test_conflicting_observations_are_not_silently_cherry_picked():
    r=rows();r.append({**r[1],'price':800})
    with pytest.raises(ValueError,match='Conflicting'):compare_pair(pair(),r,{})


@pytest.mark.parametrize('price',[0,-1,float('nan'),True,None])
def test_nonpositive_or_unknown_prices_cannot_create_ratios(price):
    r=rows();r[1]['price']=price
    with pytest.raises(ValueError,match='observed price'):compare_pair(pair(),r,{})


def test_downloaded_relative_image_path_and_resumed_absolute_path_share_cache(tmp_path,monkeypatch):
    from app.sourcing import image_provenance
    monkeypatch.chdir(tmp_path)
    config=tmp_path/'c.json'
    config.write_text(json.dumps({'products':[dict(id='oil',query_ko='oil',query_zh='oil',query_en='oil',domestic_image_indices=[0])]}))
    monkeypatch.setattr('app.sourcing_pipeline.browser_ready',lambda _:dict(ready=True))
    def fetch(url,target):
        target=target.with_suffix('.png');target.write_bytes(bytes.fromhex('89504e470d0a1a0a'))
        return target,image_provenance(target)
    monkeypatch.setattr('app.sourcing_pipeline.fetch_reference',fetch)
    calls=[]
    def run(argv,**kwargs):
        calls.append(argv)
        n='naver-shopping' in argv
        row=dict(product_id='123',item_url='https://cr3.shopping.naver.com/v2/bridge/searchGate?nv_mid=123' if n else 'https://detail.1688.com/offer/123.html' if '1688' in argv else 'https://www.alibaba.com/product-detail/Test_123.html',source_url='https://search.naver.com/search.naver?query=oil',price_min=100,currency='KRW',image_url='https://shopping-phinf.pstatic.net/query.jpg')
        if 'image-search' in argv:
            row.update(search_type='image',query_image_sha256=image_provenance(argv[argv.index('image-search')+1])['sha256'])
        return SimpleNamespace(returncode=0,stdout=json.dumps({'ok':True,'data':[row]}))
    monkeypatch.setattr('app.sourcing_pipeline.subprocess.run',run)
    assert run_pipeline(config,__import__('pathlib').Path('out'),['opencli'])['complete']
    assert len(calls)==5
    assert run_pipeline(config,__import__('pathlib').Path('out'),['opencli'],resume=True)['complete']
    assert len(calls)==5
