"""Enrich saved public catalog evidence without recollecting successful requests."""
import argparse,collections,datetime as dt,gzip,hashlib,html,json,re
from pathlib import Path
from urllib.parse import urlparse,parse_qs
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill
from app.coupang_catalog_routes import parse_ranking

def load(p):return json.loads(p.read_text())
def write(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2))
def key(r):
 q=parse_qs(urlparse(r['url']).query)
 return (str(r['product_id']),str((q.get('itemId') or [''])[0]),str((q.get('vendorItemId') or [''])[0]))
def safe(v):
 if isinstance(v,(list,dict)):v=json.dumps(v,ensure_ascii=False)
 return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v

def overseas_details(r):
 lines=[x.strip() for x in r.get('evidence_text','').splitlines() if x.strip()]
 company=next((x for x in lines if re.search(r'Co\.,?\s*Ltd|Factory|Company|Trading|株式会社|有限公司|工厂',x,re.I)),None)
 # These display signals are not audited supplier quality or monthly sales.
 rep=next((x for x in lines if re.search('재구매율|回头率',x)),None)
 years=next((x for x in lines if re.fullmatch(r'\d+\s*(년|年)',x)),None)
 score=lines[-1] if lines and re.fullmatch(r'[0-5]\.\d',lines[-1]) else None
 purchases=next((x for x in reversed(lines) if re.search(r'구매|人付款|已售|成交',x) and re.search(r'\d',x)),None)
 dispatch=next((x for x in lines if re.search(r'시간 이내 발송|小时.*发',x)),None)
 return dict(r,seller_name_observed=company,seller_years_text=years,repurchase_rate_text=rep,listing_score_untyped=float(score) if score else None,purchase_signal_text=purchases,dispatch_text=dispatch,signal_limit='listing display; score entity and purchase period not established; not monthly sales')


def main():
 a=argparse.ArgumentParser();a.add_argument('--source-root',type=Path,required=True);a.add_argument('--comparison',type=Path,required=True);a.add_argument('--out',type=Path,required=True);args=a.parse_args();O=args.out;O.mkdir(parents=True,exist_ok=True)
 sources=[];by_source={};catalog={};observations=0;checks=0
 folders=[args.source_root/n for n in ['categories','priority-categories','extended-categories','deep-categories','fashion-categories']]
 if (O/'live').exists():folders.append(O/'live')
 with (O/'observations.jsonl').open('w') as stream:
  for folder in folders:
   for f in sorted(folder.glob('*.json')):
    meta=load(f)
    if not meta.get('source_url') or not meta.get('body_sha256'):continue
    raw=f.with_suffix('.html');packed=f.with_suffix('.html.gz');body=raw.read_bytes() if raw.exists() else gzip.decompress(packed.read_bytes())
    assert hashlib.sha256(body).hexdigest()==meta['body_sha256'],f
    rows=parse_ranking(body.decode(),meta['source_url']);old={key(r):r for r in meta['data']}
    for row in rows:
     assert row['price']==old[key(row)]['price'],(f,row['product_id'])
     row.update(observed_at_kst=meta.get('observed_at_kst'),source_file=str(f),source_body_sha256=meta['body_sha256'])
     by_source[(row['source_url'],*key(row))]=row
     existing=catalog.get(row['product_id'])
     if not existing or (row.get('observed_at_kst') or '')>(existing.get('observed_at_kst') or ''):catalog[row['product_id']]=row
     stream.write(json.dumps(row,ensure_ascii=False)+'\n');observations+=1;checks+=1
    sources.append({'metadata_file':str(f),'metadata_sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'html_sha256':meta['body_sha256'],'products':len(rows)})
    if len(sources)%50==0:print(json.dumps({'stage':'reparse','sources':len(sources),'observations':observations}),flush=True)
 groups=load(args.comparison.parent.parent/'winter-products-20261023/config.json')['groups'];candidates=[];overseas=[];unmatched=[];ui={}
 for g in groups:
  folder=args.comparison/g['id'];d=load(folder/'search-response.json');q=load(folder/'qwen-selection.json');selected=set(q['selected_ids'])
  for old in d['data']:
   source_key=(old['source_url'],*key(old));row=by_source.get(source_key)
   if row is None or row['price']!=old['price']:
    unmatched.append({'group':g['id'],'id':old['product_id'],'reason':'exact source/variant/price did not match'});continue
   r=dict(row,group_id=g['id'],group_name=g['label'],application_limit=old.get('application_limit'));candidates.append(r)
   k=key(r)
   if k not in ui:ui[k]=dict(r,groups=[])
   ui[k]['groups'].append(g['label'])
  for r in load(folder/'1688-search.json')['data']:
   overseas.append(dict(overseas_details(r),group_id=g['id'],group_name=g['label'],qwen_selected=str(r.get('offer_id') or r.get('product_id')) in selected))
 assert not unmatched,unmatched
 write(O/'source-manifest.json',sources);write(O/'enriched-candidates.json',candidates);write(O/'overseas-candidates.json',overseas)
 with (O/'catalog.jsonl').open('w') as h:
  for r in catalog.values():h.write(json.dumps(r,ensure_ascii=False)+'\n')
 # Existing pair comparisons remain attached; add observed listing fields beside them.
 wb=load_workbook(args.comparison/'research.xlsx');sheet=wb['쿠팡 후보'];add=[('별점','rating'),('별점·리뷰 집계수','review_count'),('배송 배지','delivery_badges'),('로켓 배지 관찰','rocket_badge_present'),('로켓 유형','rocket_delivery_types'),('정가 표시','base_price'),('할인율 %','discount_rate_percent'),('도착 안내','delivery_text'),('무료배송 명시','free_shipping_observed'),('배송비 문구','shipping_text'),('적립 안내','cashback_text'),('품절 상태','sold_out'),('반품/중고 옵션','used_or_returned_options_text'),('반품/중고 최저가','used_or_returned_lowest_price'),('단위 가격 문구','unit_price_text'),('itemId','item_id'),('vendorItemId','vendor_item_id'),('관측 시각','observed_at_kst'),('수집 카테고리','source_url'),('판매자명','seller_name'),('적용 한계','application_limit')]
 lookup={(r['group_name'],r['product_id']):r for r in candidates};start=sheet.max_column+1
 for col,(label,_) in enumerate(add,start):sheet.cell(1,col,label)
 for row in range(2,sheet.max_row+1):
  r=lookup[(sheet.cell(row,1).value,str(sheet.cell(row,2).value))]
  for col,(_,field) in enumerate(add,start):sheet.cell(row,col,safe(r.get(field)))
 os=wb['1688 후보'];add_os=[('판매자 표시명','seller_name_observed'),('판매자 업력 표시','seller_years_text'),('재구매율 표시','repurchase_rate_text'),('표시 점수(평가 대상 미확인)','listing_score_untyped'),('구매 신호(기간 미확인)','purchase_signal_text'),('발송 안내','dispatch_text'),('관측 시각','fetched_at')];start=os.max_column+1;om={(r['group_name'],r['product_id']):r for r in overseas}
 for col,(label,_) in enumerate(add_os,start):os.cell(1,col,label)
 for row in range(2,os.max_row+1):
  r=om[(os.cell(row,1).value,str(os.cell(row,2).value))]
  for col,(_,field) in enumerate(add_os,start):os.cell(row,col,safe(r.get(field)))
 full=wb.create_sheet('쿠팡 전체 카탈로그');cols=[('ID','product_id'),('상품명','title'),('가격 KRW','price'),('별점','rating'),('집계수','review_count'),('원문','url'),*add[:-1]];full.append([x[0] for x in cols])
 for r in catalog.values():full.append([safe(r.get(k)) for _,k in cols])
 for ws in wb:
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='254B5E')
  for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=23
 bad=wb.create_sheet('원문 이상값');bad.append(['상품 ID','원본 별점','처리','목록 원문'])
 for r in catalog.values():
  if r.get('rating_quality'):bad.append([r['product_id'],r.get('rating_raw'),'0~5 범위 밖: 유효 별점에서 제외',r['source_url']])
 wb.save(O/'research.xlsx')
 summary={'unique_catalog_products':len(catalog),'winter_unique_variants':len(ui),'winter_group_observations':len(candidates),'overseas_observations':len(overseas),'source_pages':len(sources),'all_source_observations':observations,'winter_rating_available':sum(r['rating'] is not None for r in candidates),'winter_review_count_available':sum(r['review_count'] is not None for r in candidates),'winter_rocket_badged':sum(r['rocket_badge_present'] for r in candidates),'winter_delivery_available':sum(bool(r['delivery_text']) for r in candidates),'winter_free_shipping_explicit':sum(r['free_shipping_observed'] is True for r in candidates),'winter_seller_name_available':sum(r['seller_name'] is not None for r in candidates),'overseas_seller_name_available':sum(bool(r['seller_name_observed']) for r in overseas),'overseas_score_available':sum(r['listing_score_untyped'] is not None for r in overseas),'price_invariance_checks':checks,'saved_at_kst':dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).isoformat(),'fresh_live_sources':len(list((O/'live').glob('*.json'))) if (O/'live').exists() else 0,'price_basis':'listing-time, variant-specific; not landed cost','monthly_sales_verified':False}
 # Compact UI data excludes verbose card text but keeps exact source links.
 fields=['rating_raw','rating_quality','product_id','title','price','rating','review_count','image_url','url','delivery_badges','rocket_badge_present','rocket_delivery_types','base_price','discount_rate_percent','delivery_text','free_shipping_observed','shipping_text','cashback_text','sold_out','unit_price_text','used_or_returned_options_text','used_or_returned_lowest_price','item_id','vendor_item_id','observed_at_kst','source_url','groups','application_limit']
 winter=[{k:r.get(k) for k in fields} for r in ui.values()];allrows=[{k:r.get(k) for k in fields} for r in catalog.values()]
 payload=json.dumps({'winter':winter,'catalog':allrows,'overseas':overseas,'summary':summary},ensure_ascii=False,separators=(',',':')).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
 template=Path(__file__).with_name('coupang_catalog_enrich.html').read_text();(O/'index.html').write_text(template.replace('__DATA__',payload))
 write(O/'summary.json',summary);print(json.dumps(summary,ensure_ascii=False),flush=True)
if __name__=='__main__':main()
