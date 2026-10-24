import copy
import json
from pathlib import Path
from decimal import Decimal
import pytest
from app.rocket_margin import calculate, minimum_price


def inputs():
    cfg=json.loads((Path(__file__).parents[1]/'configs/rocket-margin.example.json').read_text())
    a={**cfg['defaults'],**cfg['scenarios']['base']}
    # A hand-checkable taxable importer example with no operational reserves.
    for k in list(a):
        if isinstance(a[k],(int,float)): a[k]=0
    a.pop('shipment', None)
    a.update(sale_factor=1,fee_rate=.1,batch_units=100,logistics_multiplier=1,fee_rounding='ceil_10')
    item=dict(domestic={'price':11000},overseas={'price_min':1000},quantity_ratio_assumed=1,review={'risk_traits':[]})
    return item,a,dict(io=100,delivery=200)


def test_independent_cash_and_accrual_ledgers():
    i,a,l=inputs();r=calculate(i,a,l)
    assert r['revenue_ex_vat']==10000
    assert r['output_vat']==1000
    assert r['input_vat_deductible']==240
    assert r['cash_contribution_before_vat']==8360
    assert r['vat_due_allocated']==760
    assert r['operating_profit_after_vat']==7600
    assert r['operating_profit_after_vat']==r['revenue_ex_vat']-sum(x['profit_cost'] for x in r['ledger'])


def test_foreign_purchase_vat_is_not_imaginary_domestic_credit():
    i,a,l=inputs();r=calculate(i,a,l)
    foreign=next(x for x in r['ledger'] if x['name']=='overseas_goods')
    assert foreign['vat']==0
    assert r['import_vat']==100
    assert next(x for x in r['ledger'] if x['name']=='import_vat')['deductible_vat']==100
    no_credit=calculate(i,a,l,credit_ratio=0)
    assert no_credit['operating_profit_after_vat']==7360
    assert no_credit['input_vat_deductible']==0


def test_coupon_reduces_revenue_fee_and_output_vat_once():
    i,a,l=inputs();a['coupon']=1100;r=calculate(i,a,l)
    assert r['final_consumer_price']==9900
    assert r['output_vat']==900
    assert next(x for x in r['ledger'] if x['name']=='commission')['net']==990
    assert r['operating_profit_after_vat']==6710


def test_breakeven_is_minimum_integer_after_coupon_price():
    i,a,l=inputs();a.update(ads_rate=.15,coupon=500)
    p=minimum_price(i,a,l)
    assert calculate(i,a,l,consumer_price=p)['operating_profit_after_vat']>=0
    assert calculate(i,a,l,consumer_price=p-1)['operating_profit_after_vat']<0
    goal=minimum_price(i,a,l,target_margin=.2)
    r=calculate(i,a,l,consumer_price=goal)
    assert r['operating_profit_after_vat']>=r['revenue_ex_vat']*Decimal('.2')


def test_breakeven_ad_budget_prices_vat_consistently():
    i,a,l=inputs()
    for credit in (0,1):
        r=calculate(i,a,l,credit_ratio=credit)
        a2=copy.deepcopy(a)
        a2['ads_rate']=r['max_ad_spend_ex_vat']/r['final_consumer_price']
        near=calculate(i,a2,l,credit_ratio=credit)
        assert abs(near['operating_profit_after_vat'])<Decimal('.02')


def test_returns_add_outbound_handling_and_damage_without_double_refund():
    i,a,l=inputs();base=calculate(i,a,l)
    a.update(return_rate=.1,return_collection=100,return_restock=50,return_writeoff_fraction=.5)
    r=calculate(i,a,l)
    assert r['final_consumer_price']==base['final_consumer_price']
    assert r['operating_profit_after_vat']<base['operating_profit_after_vat']
    assert next(x for x in r['ledger'] if x['name']=='returned_order_outbound')['net']==Decimal('33.33')
    assert next(x for x in r['ledger'] if x['name']=='return_damage_reserve')['net']>0


def test_inventory_and_cash_recovery_require_seller_forecast():
    i,a,l=inputs();r=calculate(i,a,l)
    assert r['initial_inventory_cash']==110000
    assert r['monthly_operating_profit'] is None
    assert r['inventory_days'] is None
    assert r['capital_recovery_days_excluding_settlement'] is None
    a.update(monthly_seller_units=100,settlement_delay_days=14)
    r=calculate(i,a,l)
    assert r['inventory_days']==30
    assert r['capital_recovery_days_excluding_settlement']<30
    assert r['capital_recovery_days_including_settlement']==r['capital_recovery_days_excluding_settlement']+14


def test_storage_promotion_requires_explicit_days():
    i,a,l=inputs();a.update(held_days=45,storage_free_days=30,storage_daily=5)
    r=calculate(i,a,l)
    assert next(x for x in r['ledger'] if x['name']=='storage')['net']==75
    a['storage_free_days']=60;r=calculate(i,a,l)
    assert next(x for x in r['ledger'] if x['name']=='storage')['net']==0


def test_default_purchase_cost_is_observed_1688_price_without_markup():
    cfg=json.loads((Path(__file__).parents[1]/'configs/rocket-margin.example.json').read_text())
    item,_,logistics=inputs()
    for scenario in cfg['scenarios'].values():
        a={**cfg['defaults'],**scenario}
        assert a['purchase_buffer']==0
        result=calculate(item,a,logistics)
        goods=next(x for x in result['ledger'] if x['name']=='overseas_goods')
        assert goods['net']==item['overseas']['price_min']
        assert goods['vat']==0
        assert result['customs_value_proxy']==Decimal(1000)+Decimal(str(a['shipment']['totals']['china_inland']))/100+Decimal(str(a['shipment']['totals']['international']))/100


@pytest.mark.parametrize('key,value',[('return_rate',1),('unsold_fraction',1),('duty_rate',-1),('ads_rate',2),('china_inland',float('nan')),('coupon',12000)])
def test_invalid_inputs_do_not_silently_become_zero(key,value):
    i,a,l=inputs();a[key]=value
    with pytest.raises(ValueError):calculate(i,a,l)


def test_report_roundtrip_preserves_missing_seller_forecast_and_escapes_html(tmp_path):
    from app.rocket_margin_report import build
    from openpyxl import load_workbook
    i,a,l=inputs()
    i.update(pair_id='test:1',quantity_verified=False,eligible=True,
             scenarios={'base':{'estimated_profit':7777}})
    i['domestic'].update(title='<script>not executable</script>',group_name='시험',url='https://www.coupang.com/vp/products/1')
    i['overseas'].update(title='원가 미확정',item_url='https://detail.1688.com/offer/1.html')
    i['review']['reason_ko']='구성 확인 필요'
    i['monthly_review_metrics']={'estimated_monthly_sales':1500,'status':'complete'}
    source=tmp_path/'input.json';source.write_text(json.dumps({'items':[i]}))
    config=tmp_path/'config.json';config.write_text(json.dumps({'defaults':a,'scenarios':{'base':{}},'size_starting_fees_ex_vat':{'중형':l}}))
    out=tmp_path/'out';summary=build(source,config,out)
    data=json.loads((out/'report.json').read_text())
    assert summary['pairs']==1
    assert data['records'][0]['monthly_operating_profit'] is None
    assert data['products'][0]['catalog_monthly_sales']==1500
    assert data['records'][0]['operating_profit_after_vat']==7600
    assert '<script>not executable</script>' not in (out/'index.html').read_text()
    book=load_workbook(out/'로켓그로스_마진분석.xlsx',data_only=True)
    assert book['전체 마진 민감도'].max_row==2
    assert book['전체 마진 민감도'].cell(2,10).value==7600
    assert book['전체 마진 민감도'].cell(2,21).value is None


def test_mixed_batch_allocation_preserves_total_and_uses_measured_weights():
    from app.rocket_margin import allocate_batch
    items=[dict(id='a',quantity=10,unit_weight=2,unit_volume=3,unit_value=5),
           dict(id='b',quantity=20,unit_weight=4,unit_volume=1,unit_value=10)]
    assert allocate_batch(1000,items,'weight')=={'a':200,'b':800}
    assert allocate_batch(1000,items,'volume')=={'a':600,'b':400}
    assert allocate_batch(1000,items,'value')=={'a':200,'b':800}
    tiny=allocate_batch('.02',[dict(id=str(n),quantity=1) for n in range(4)])
    assert sum(tiny.values())==Decimal('.02') and min(tiny.values())>=0
    with pytest.raises(KeyError):allocate_batch(1000,[dict(id='a',quantity=1)],'weight')
    with pytest.raises(ValueError):allocate_batch(1000,[dict(id='a',quantity=0)])


def test_consolidated_freight_quantity_ratio_and_no_bulky_guess():
    from app.rocket_margin import BATCH_COSTS
    i,a,l=inputs();i['quantity_ratio_assumed']=2
    a.update(batch_units=250,shipment={'basis':'units','totals':dict.fromkeys(BATCH_COSTS,0)})
    a['shipment']['totals']['international']=100000
    r=calculate(i,a,l)
    # 500 supplier pieces = 250 domestic two-piece selling units.
    assert r['shipment_allocation']['shipment_physical_units']==500
    assert r['shipment_allocation']['per_sale_costs']['international']==400
    assert next(x for x in r['ledger'] if x['name']=='international_freight')['net']==400
    i['review']['risk_traits']=['bulky']
    assert calculate(i,a,l)['operating_profit_after_vat']==r['operating_profit_after_vat']
    a['batch_units']=500
    assert calculate(i,a,l)['shipment_allocation']['per_sale_costs']['international']==200
    a['international']=100
    with pytest.raises(ValueError):calculate(i,a,l)


def test_mixed_shipment_derives_actual_sku_inventory_and_cost():
    from app.rocket_margin import BATCH_COSTS
    i,a,l=inputs();i['pair_id']='a';i['quantity_ratio_assumed']=2
    a['shipment']={'basis':'weight','totals':dict.fromkeys(BATCH_COSTS,0),
        'items':[dict(id='a',quantity=100,unit_weight=2),dict(id='b',quantity=200,unit_weight=4)]}
    a['shipment']['totals']['international']=100000
    r=calculate(i,a,l)
    assert r['batch_units']==50
    assert r['shipment_allocation']['sku_allocated_totals']['international']==20000
    assert r['shipment_allocation']['per_sale_costs']['international']==400
