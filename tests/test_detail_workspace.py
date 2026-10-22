import io
import json
import threading
from http.server import ThreadingHTTPServer
import httpx
from openpyxl import Workbook, load_workbook
import pytest
from app.detail_workspace import Store, Conflict
from app.detail_workspace_web import handler

@pytest.fixture
def store(tmp_path):return Store(tmp_path)

def setup_page(s):
    pid=s.add_product({'name':'노란 부적','source_url':'https://example.com/p'})
    p=s.add_page(pid,{'title':'A안','html':'<h1>처음</h1>'})
    return pid,p

def test_instruction_apply_persistence_and_conflict(store):
    pid,p=setup_page(store);p=store.save_page(p['id'],{'revision':1,'plan':'타깃과 핵심 약속','html':'<h1>사람 수정</h1>'})
    store.request(p['id'],{'base_revision':2,'body':'제목을 단정하게','target':'첫 화면'})
    h=Store(store.root).handoff(p['id']);assert h['page']['plan']=='타깃과 핵심 약속';q=h['instructions'][0]
    with pytest.raises(Conflict):store.apply(p['id'],{'base_revision':1,'request_ids':[q['id']],'html':'틀린 버전'})
    result=store.apply(p['id'],{'base_revision':2,'request_ids':[q['id']],'html':'<h1>새 제목</h1>'});assert result['revision']==3;assert result['status']=='검토 중';assert result['requests'][0]['applied_revision']==3
    with pytest.raises(ValueError):store.apply(p['id'],{'base_revision':3,'request_ids':[q['id']],'html':'중복 적용'})
    restored=store.restore(p['id'],{'revision':3,'target_revision':1});assert restored['revision']==4;assert restored['html']=='<h1>처음</h1>';assert len(restored['history'])==4

def test_batch_atomic_and_url_validation(store):
    with pytest.raises(ValueError):store.import_rows([{'name':'첫 상품'},{'name':'잘못된 상품','source_url':'javascript:alert(1)'}])
    assert store.state()['products']==[]
    with pytest.raises(ValueError):store.add_product({'name':'잘못된 링크','source_url':'javascript:alert(1)'})
    _,p=setup_page(store)
    with pytest.raises(ValueError):store.add_page(p['product_id'],{'title':'나쁜 편집기','editor_url':'file:///etc/passwd'})

def test_excel_header_mapping_and_export(store):
    w=Workbook();ws=w.active;ws.append(['조사 목록']);ws.append(['이름 사용자 열','판매 주소','담당']);ws.append(['부적','https://example.com','담당A']);b=io.BytesIO();w.save(b)
    v=store.import_excel(b.getvalue(),header_row=2,mapping={'name':0,'source_url':1,'owner':2});assert v['rows'][0]['owner']=='담당A';store.import_rows(v['rows'])
    _,p=setup_page(store);store.request(p['id'],{'base_revision':1,'body':'수정 내용'})
    out=load_workbook(io.BytesIO(store.export()));assert out['상품 목록'].max_row==3;assert out['상세페이지 목록'].max_row==2;assert out['수정 지시'].max_row==2

def test_loopback_csrf_and_preview_script_policy(store):
    _,p=setup_page(store);server=ThreadingHTTPServer(('127.0.0.1',0),handler(store,'secret'));t=threading.Thread(target=server.serve_forever,daemon=True);t.start();base=f'http://127.0.0.1:{server.server_port}'
    try:
        with httpx.Client(base_url=base,trust_env=False) as c:
            assert c.post('/api/products',json={'name':'CSRF'}).status_code==403
            assert c.get('/api/state',headers={'Host':'evil.example'}).status_code==403
            r=c.post('/api/products',json={'name':'ok'},headers={'X-CSRF-Token':'secret'});assert r.status_code==201
            assert c.get('/api/pages/'+p['id']+'/handoff').json()['page']['html']=='<h1>처음</h1>'
            assert "script-src 'none'" in c.get('/pages/'+p['id']+'/preview').headers['Content-Security-Policy']
    finally:server.shutdown();server.server_close();t.join()

def test_production_registration_preserves_edits_and_links(store,tmp_path):
    from PIL import Image
    from app.detail_production import register
    image=tmp_path/'photo.png';Image.new('RGB',(80,160),'yellow').save(image)
    report=tmp_path/'source.json';report.write_text(json.dumps({'candidates':[{'key':'1688/'+key,'item_url':'https://detail.1688.com/offer/'+key+'.html','photo':{'path':str(image)},'price_text':'¥0.13-0.22','moq_text':'500个'} for key in ['973958651038','1073499594799']]}))
    first=register(store,report,'http://127.0.0.1:9876');assert len(first)==2
    page=store.get_page(first[0]['page_id']);assert page['kind']=='draft';assert '확정할 항목' in page['plan'];assert 'id="hero"' in page['html'];store.save_page(page['id'],{'revision':1,'html':'<h1>사람 수정 보존</h1>'})
    second=register(store,report);assert [x['page_id'] for x in first]==[x['page_id'] for x in second];assert len(store.state()['products'])==2;assert store.get_page(page['id'])['html']=='<h1>사람 수정 보존</h1>'
    out=load_workbook(io.BytesIO(store.export('http://127.0.0.1:9876')));sheet=out['상세페이지 목록'];assert sheet['G2'].hyperlink.target.startswith('http://127.0.0.1:9876/?page=');assert sheet['J2'].hyperlink.target.endswith('/preview');assert sheet['K2'].hyperlink.target.endswith('/handoff')
